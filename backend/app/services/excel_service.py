import io
from datetime import date, datetime

from openpyxl import Workbook, load_workbook
from sqlalchemy.orm import Session

from ..models import Cabinet, ChangeLog, Equipment, Room
from .spec_utils import cell_to_specs, specs_to_cell

# 导出表头：前 15 列沿用原顺序（第 10 列由 IP 升级为带内IPv4），末尾追加新列
HEADERS = [
    "资产编号", "设备名称", "类型", "品牌", "型号", "SN序列号",
    "机房", "机柜", "U位", "带内IPv4", "状态", "采购日期", "保修到期",
    "责任人", "备注",
    "带内IPv6", "带外IPv4", "带外IPv6", "CPU", "显卡", "硬盘",
]

# 兼容旧模板：旧文件的 IP 列映射到带内IPv4
HEADER_ALIASES = {"IP": "带内IPv4"}

VALID_STATUS = ("在用", "备用", "维修", "报废", "退役")
VALID_CATEGORY = ("服务器", "交换机", "存储", "防火墙", "其他")

# 字段定义：(模型字段, 表头, 类型)
TEXT_FIELDS = [
    ("name", "设备名称", "text"),
    ("category", "类型", "text"),
    ("brand", "品牌", "text"),
    ("model", "型号", "text"),
    ("sn", "SN序列号", "text"),
    ("u_position", "U位", "text"),
    ("status", "状态", "text"),
    ("purchase_date", "采购日期", "date"),
    ("warranty_end", "保修到期", "date"),
    ("owner", "责任人", "text"),
    ("remark", "备注", "text"),
]
IP_FIELDS = [
    ("ip_inband_v4", "带内IPv4"),
    ("ip_inband_v6", "带内IPv6"),
    ("ip_outband_v4", "带外IPv4"),
    ("ip_outband_v6", "带外IPv6"),
]
SPEC_FIELDS = [
    ("cpus", "CPU", "cpu"),
    ("gpus", "显卡", "gpu"),
    ("disks", "硬盘", "disk"),
]


def generate_template() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "设备台账"
    ws.append(HEADERS)
    ws.append(
        ["DC-2026-0001", "示例服务器", "服务器", "Dell", "R740",
         "SN123456", "机房A", "A-01", "U10", "10.0.0.10", "在用",
         "2025-01-15", "2028-01-15", "张三", "示例数据，导入时会自动更新该行",
         "", "10.0.0.11", "",
         '[{"model": "Intel Xeon Gold 6338", "sockets": 2, "cores_per_cpu": 32, "freq_ghz": 2.0}]',
         '[{"model": "NVIDIA A100", "count": 4, "memory_gb": 80, "purpose": "计算"}]',
         '[{"type": "NVMe", "capacity_gb": 3840, "count": 8, "raid_level": "RAID10", "role": "数据"}]']
    )
    widths = [14, 16, 10, 10, 16, 16, 10, 10, 8, 14, 8, 12, 12, 10, 20,
              14, 14, 14, 40, 32, 40]
    for col, width in zip("ABCDEFGHIJKLMNOPQRSTU", widths):
        ws.column_dimensions[col].width = width
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_equipment(rows: list[Equipment]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "设备台账"
    ws.append(HEADERS)
    for e in rows:
        ws.append(
            [
                e.asset_no, e.name, e.category, e.brand, e.model, e.sn,
                e.room.name if e.room else None,
                e.cabinet.name if e.cabinet else None,
                e.u_position, e.ip_inband_v4, e.status,
                e.purchase_date.isoformat() if e.purchase_date else None,
                e.warranty_end.isoformat() if e.warranty_end else None,
                e.owner, e.remark,
                e.ip_inband_v6, e.ip_outband_v4, e.ip_outband_v6,
                specs_to_cell(e.cpus), specs_to_cell(e.gpus), specs_to_cell(e.disks),
            ]
        )
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _parse_date(value) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    s = str(value).strip()
    for fmt in ("%Y-%m-%d", "%Y-%m-%d %H:%M:%S", "%Y/%m/%d", "%Y/%m/%d %H:%M:%S"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"无法识别的日期格式：{s}")


def _get_or_create_room(db: Session, name: str | None) -> Room | None:
    if not name:
        return None
    name = str(name).strip()
    room = db.query(Room).filter(Room.name == name).first()
    if not room:
        room = Room(name=name)
        db.add(room)
        db.flush()
    return room


def _find_cabinet(db: Session, room: Room | None, name) -> Cabinet | None:
    if not room or name in (None, ""):
        return None
    return (
        db.query(Cabinet)
        .filter(Cabinet.room_id == room.id, Cabinet.name == str(name).strip())
        .first()
    )


def _build_header_map(header_row) -> dict:
    header_map = {}
    for idx, h in enumerate(header_row):
        if h is None:
            continue
        name = str(h).strip()
        header_map[HEADER_ALIASES.get(name, name)] = idx
    return header_map


def import_equipment(db: Session, content: bytes, operator_id: int | None) -> dict:
    wb = load_workbook(io.BytesIO(content), data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)

    try:
        header_row = next(rows)
    except StopIteration:
        return {"created": 0, "updated": 0, "failed": 0, "errors": [], "warnings": []}
    header_map = _build_header_map(header_row)

    created = updated = failed = 0
    errors: list[str] = []
    warnings: list[str] = []

    def cell(row, header):
        i = header_map.get(header)
        if i is None or i >= len(row) or row[i] is None:
            return None
        return str(row[i]).strip()

    for idx, row in enumerate(rows, start=2):
        asset_no = cell(row, "资产编号")
        if not asset_no:
            continue  # 空行/无资产编号，跳过

        try:
            existing = (
                db.query(Equipment).filter(Equipment.asset_no == asset_no).first()
            )

            category = cell(row, "类型")
            status = cell(row, "状态")
            name = cell(row, "设备名称")

            if category is not None and category not in VALID_CATEGORY:
                raise ValueError(f"类型必须是 {'/'.join(VALID_CATEGORY)}")
            if status is not None and status not in VALID_STATUS:
                raise ValueError(f"状态必须是 {'/'.join(VALID_STATUS)}")
            if existing is None:
                if not name:
                    raise ValueError("设备名称为必填项")
                if not category:
                    raise ValueError("类型为必填项")

            # 只收集“已填写”的字段：更新时空单元格保留原值，新增时缺失字段走模型默认
            data: dict = {}
            for key, header, kind in TEXT_FIELDS:
                val = cell(row, header)
                if val is None:
                    continue
                data[key] = _parse_date(val) if kind == "date" else val

            for key, header in IP_FIELDS:
                val = cell(row, header)
                if val is not None:
                    data[key] = val

            for key, header, kind in SPEC_FIELDS:
                val = cell(row, header)
                specs, warn = cell_to_specs(val, kind)
                if warn:
                    warnings.append(f"第 {idx} 行：{warn}")
                if specs is not None:
                    data[key] = specs

            # 位置：机房/机柜
            room_name = cell(row, "机房")
            cabinet_name = cell(row, "机柜")
            if room_name is not None:
                room = _get_or_create_room(db, room_name)
                cabinet = _find_cabinet(db, room, cabinet_name)
                data["room_id"] = room.id if room else None
                data["cabinet_id"] = cabinet.id if cabinet else None
            elif cabinet_name is not None and existing is not None:
                cabinet = _find_cabinet(db, existing.room, cabinet_name)
                data["cabinet_id"] = cabinet.id if cabinet else None

            # SN 唯一性
            sn = data.get("sn")
            if sn:
                q = db.query(Equipment).filter(Equipment.sn == sn)
                if existing is not None:
                    q = q.filter(Equipment.id != existing.id)
                if q.first():
                    raise ValueError(f"SN 序列号 {sn} 已被其他设备使用")

            if existing is not None:
                for k, v in data.items():
                    setattr(existing, k, v)
                created_flag = False
            else:
                existing = Equipment(asset_no=asset_no, **data)
                db.add(existing)
                created_flag = True

            db.flush()
            db.add(
                ChangeLog(
                    equipment_id=existing.id,
                    asset_no=asset_no,
                    action="导入",
                    detail={"created": created_flag},
                    operator_id=operator_id,
                )
            )
            if created_flag:
                created += 1
            else:
                updated += 1
        except Exception as exc:  # noqa: BLE001 - 逐行收集错误继续导入
            failed += 1
            errors.append(f"第 {idx} 行：{exc}")

    db.commit()
    return {
        "created": created,
        "updated": updated,
        "failed": failed,
        "errors": errors,
        "warnings": warnings,
    }
